"""Extract tasks from .txt / .md / .docx / .pdf / .xlsx / .csv documents.

Two layouts are understood:

Table layout (checked first): the document has tables with a "Date" and "Task" column,
e.g. a PDF or Word report with one table per work area.
  * A numbered heading such as "1. AWS Integration & Dashboard" (or a Markdown/Word heading)
    above a table sets the category for its rows; a table at the top of a page with no heading
    continues the previous category.
  * Optional columns: Category, Status, Type.
  * A summary table with "Category" and "No. of Tasks" columns is not imported, but its
    counts are returned so the app can confirm every task was found.

Line layout (everything else):
  * A line that is (or starts with) a date becomes the current date context.
  * Bullet lines ("->", "-", "*", "•", "1.", "1)") become tasks.
  * If the document has no bullet lines at all, every other non-empty line is a task.
  * Separator lines (====, ----) and empty lines are ignored.

In both layouts each task gets
  * a kind: "bug", "feature" or "task", detected (in priority order) from
      1. an explicit tag on the task:  "[Bug] ...", "(feature) ...", "Bug: ...", "Task - ...", "... #bug"
         (or a Type column)
      2. a section heading above it:  "Bugs:", "Features", "## Tasks"
      3. a keyword in the text:        "bug", "defect", "hotfix" / "feature", "enhancement", "story" / "task"
    Tasks with no kind use the default issue type chosen in the app.
  * a status hint from a Status column, a tag ("[done]", "(wip)", "[pending]") or a phrase
    ("in progress", "incomplete", "completed", "moved to prod"). The app maps it to a Jira status.
  * a leave flag for rows such as "On Leave", so they are not created by default.
"""
import csv
import io
import re
from collections import Counter
from datetime import date, datetime

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

DASHES = "–—"  # en dash, em dash
DATE_NUMERIC = re.compile(r"(\d{1,4})[/\-.](\d{1,2})[/\-.](\d{1,4})")
DATE_TEXT_1 = re.compile(r"(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]{3,9})\.?,?\s+(\d{4})")   # 1 Sep 2026
DATE_TEXT_2 = re.compile(r"([A-Za-z]{3,9})\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})")   # Sep 1, 2026
DATE_LINE = re.compile(rf"^\s*(?:date\s*[:\-{DASHES}]*\s*)?(?P<d>.+?)\s*[:\-{DASHES}]*\s*$", re.I)
BULLET = re.compile(
    r"^\s*(?:->|=>|[-*•·▪◦‣]+|\d+[.)]|[a-zA-Z][.)]|\[[ xX]?\])\s+(?P<text>.+?)\s*$")
SEPARATOR = re.compile(r"^[\s=\-_*#~.]+$")

# ---------- bug / feature / task detection ----------

KIND_WORDS = {
    "bug": r"bugs?|defects?|hotfix(?:es)?|bugfix(?:es)?",
    "feature": r"new\s+features?|features?|enhancements?|stor(?:y|ies)",
    "task": r"tasks?|chores?",
}
_ANY_KIND = "|".join(KIND_WORDS.values())
_KIND_RE = {k: re.compile(rf"^(?:{v})$", re.I) for k, v in KIND_WORDS.items()}
_SEP = rf":\-{DASHES}|"
# "[Bug] x", "(bug) x", "{bug} x", "Bug: x", "Bug - x", "Bug | x"
TAG_PREFIX = re.compile(
    rf"^\s*(?:[\[({{]\s*(?P<w1>{_ANY_KIND})\s*[\])}}]\s*[{_SEP}]?"
    rf"|(?P<w2>{_ANY_KIND})\s*[{_SEP}])\s*", re.I)
# "x [bug]", "x (feature)", "x #bug"
TAG_SUFFIX = re.compile(
    rf"\s*(?:[\[({{]\s*(?P<w1>{_ANY_KIND})\s*[\])}}]|#(?P<w2>{_ANY_KIND}))\s*$", re.I)
# Heading lines such as "Bugs:", "## Features", "Tasks -"
SECTION = re.compile(rf"^\s*#*\s*(?P<w1>{_ANY_KIND})\s*[:\-{DASHES}]*\s*$", re.I)
KEYWORD = re.compile(rf"\b(?P<w1>{_ANY_KIND})\b", re.I)


def _kind_of(match: re.Match) -> str | None:
    groups = match.groupdict()
    word = re.sub(r"\s+", " ", groups.get("w1") or groups.get("w2") or "")
    return next((k for k, rx in _KIND_RE.items() if rx.match(word)), None)


def detect_kind(text: str) -> tuple[str, str | None, str | None]:
    """Return (text without tag, kind, source) where source is 'tag' or 'keyword'."""
    m = TAG_PREFIX.match(text)
    if m and m.end() < len(text):
        return text[m.end():].strip(), _kind_of(m), "tag"
    m = TAG_SUFFIX.search(text)
    if m and m.start() > 0:
        return text[:m.start()].strip(), _kind_of(m), "tag"
    kinds = {_kind_of(k) for k in KEYWORD.finditer(text)}
    # A bug mention wins ("feature X has a bug"), then feature, then task.
    for k in ("bug", "feature", "task"):
        if k in kinds:
            return text, k, "keyword"
    return text, None, None


# ---------- status hints ----------
# Jira status categories: "new" (To Do), "indeterminate" (In Progress), "done".

STATUS_TAG_WORDS = {
    "done": r"done|completed?|finished|closed|resolved",
    "indeterminate": r"in[ -]?progress|wip|ongoing|doing|started",
    "new": r"pending|to[ -]?do|not started|on hold|blocked",
}
_ANY_STATUS = "|".join(STATUS_TAG_WORDS.values())
_STATUS_RE = {k: re.compile(rf"^(?:{v})$", re.I) for k, v in STATUS_TAG_WORDS.items()}
STATUS_TAG_PREFIX = re.compile(rf"^\s*[\[(]\s*(?P<w>{_ANY_STATUS})\s*[\])]\s*[:\-{DASHES}]?\s*", re.I)
STATUS_TAG_SUFFIX = re.compile(rf"\s*[\[(]\s*(?P<w>{_ANY_STATUS})\s*[\])]\s*$", re.I)
# Phrases inside the text, checked in this order ("not completed" is in progress, not done).
STATUS_PHRASES = [
    ("indeterminate", re.compile(
        r"\b(?:in[ -]?progress|wip|ongoing|incomplete|partially|half[ -]done|started|"
        r"not (?:yet )?(?:completed?|done|finished))\b", re.I)),
    ("new", re.compile(r"\b(?:pending|not started|yet to start|on hold|blocked)\b", re.I)),
    ("done", re.compile(r"\b(?:completed?|done|finished|moved to prod(?:uction)?|deployed|released|merged)\b", re.I)),
]
LEAVE = re.compile(r"^(?:on\s+)?(?:leave|holiday|public holiday|sick leave|day off|off|vacation|pto)$", re.I)


def status_category(text: str) -> str | None:
    """Jira status category suggested by a status word or phrase."""
    text = (text or "").strip()
    for k, rx in _STATUS_RE.items():
        if rx.match(text):
            return k
    for k, rx in STATUS_PHRASES:
        if rx.search(text):
            return k
    return None


def detect_status(text: str) -> tuple[str, dict | None]:
    """Return (text without a status tag, status hint). A hint is
    {"category": "done"|"indeterminate"|"new"|None, "name": str|None, "source": "tag"|"keyword"|"column"}."""
    for rx, cut in ((STATUS_TAG_PREFIX, lambda m: text[m.end():]), (STATUS_TAG_SUFFIX, lambda m: text[:m.start()])):
        m = rx.search(text)
        if m and cut(m).strip():
            return cut(m).strip(), {"category": status_category(m.group("w")), "name": None, "source": "tag"}
    for k, rx in STATUS_PHRASES:
        if rx.search(text):
            return text, {"category": k, "name": None, "source": "keyword"}
    return text, None


def is_leave(*texts: str) -> bool:
    return any(LEAVE.match((t or "").strip()) for t in texts)


# ---------- dates ----------

def _month(name: str) -> int | None:
    return MONTHS.get(name[:3].lower())


def parse_date(text: str) -> date | None:
    """Return a date if the text is a date (dd/mm/yyyy assumed when ambiguous)."""
    text = text.strip()
    m = DATE_NUMERIC.fullmatch(text)
    if m:
        a, b, c = (int(x) for x in m.groups())
        try:
            if a > 31:                      # yyyy-mm-dd
                return date(a, b, c)
            if c < 100:
                c += 2000
            if a > 12 and b <= 12:          # dd/mm
                return date(c, b, a)
            if b > 12 and a <= 12:          # mm/dd
                return date(c, a, b)
            return date(c, b, a)            # ambiguous -> day first
        except ValueError:
            return None
    m = DATE_TEXT_1.fullmatch(text)
    if m and _month(m.group(2)):
        try:
            return date(int(m.group(3)), _month(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    m = DATE_TEXT_2.fullmatch(text)
    if m and _month(m.group(1)):
        try:
            return date(int(m.group(3)), _month(m.group(1)), int(m.group(2)))
        except ValueError:
            return None
    return None


def line_date(line: str) -> date | None:
    """Detect a date header line such as 'Date:- 01/09/2026' or '## 1 Sep 2026'."""
    stripped = line.strip().lstrip("#").strip()
    if not stripped or len(stripped) > 40:
        return None
    m = DATE_LINE.match(stripped)
    if not m:
        return None
    return parse_date(m.group("d"))


# ---------- text extraction ----------

def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        import docx
        document = docx.Document(io.BytesIO(data))
        parts = []
        for p in document.paragraphs:
            style = (p.style.name or "").lower() if p.style is not None else ""
            text = p.text
            if "list" in style and text.strip() and not BULLET.match(text):
                text = "- " + text.strip()
            parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append(" | ".join(c for c in cells if c))
        return "\n".join(parts)
    if name.endswith(".pdf"):
        import pymupdf
        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            return "\n".join(page.get_text() for page in pdf)
    if name.endswith(".doc"):
        raise ValueError("Legacy .doc files are not supported. Save it as .docx and try again.")
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# ---------- task extraction ----------

def parse_tasks(text: str) -> list[dict]:
    lines = text.splitlines()
    has_bullets = any(BULLET.match(l) for l in lines)
    tasks: list[dict] = []
    current: date | None = None
    section_kind: str | None = None
    for raw in lines:
        line = raw.rstrip()
        if not line.strip() or SEPARATOR.match(line):
            continue
        d = line_date(line)
        if d:
            current = d
            section_kind = None  # each day starts without a section
            continue
        bullet = BULLET.match(line)
        section = None if bullet else SECTION.match(line)
        if section:
            section_kind = _kind_of(section)
            continue
        if bullet:
            text_ = bullet.group("text").strip()
        elif has_bullets:
            continue  # headings / titles between bullets are not tasks
        else:
            text_ = line.strip()
        if not text_:
            continue
        tasks.append(make_task(len(tasks) + 1, text_, current, section_kind=section_kind))
    return tasks


def make_task(id_: int, text: str, when: date | None, category: str | None = None,
              section_kind: str | None = None, status_text: str = "", kind_text: str = "") -> dict:
    text, kind, source = detect_kind(text)
    if kind_text:
        m = KEYWORD.search(kind_text)
        if m and _kind_of(m):
            kind, source = _kind_of(m), "tag"
    if source != "tag" and section_kind:
        kind, source = section_kind, "section"
    text, hint = detect_status(text)
    if status_text:
        hint = {"category": status_category(status_text), "name": status_text, "source": "column"}
    return {
        "id": id_,
        "date": when.isoformat() if when else None,
        "summary": text,
        "kind": kind,
        "kind_source": source,
        "category": category or None,
        "status_hint": hint,
        "leave": is_leave(text, category or ""),
    }


# ---------- table layout ----------

# Header cell text -> column role. Matched against the whole cell, so "No. of Tasks" is not "Tasks".
COLUMN_ROLES = {
    "date": {"date", "day", "dated", "work date", "date of work"},
    "summary": {"task", "tasks", "task name", "task details", "task description", "work", "work done",
                "work item", "description", "summary", "details", "activity", "activities", "item", "title"},
    "category": {"category", "module", "area", "work area", "epic", "component", "section", "group",
                 "feature area", "parent"},
    "status": {"status", "state", "progress"},
    "kind": {"type", "issue type", "kind"},
}
COUNT_HEADER = re.compile(r"^(?:(?:no\.?|number|#|count)\s*(?:of\s+)?tasks?|tasks?\s*count|count|total)$")
NUMBERED_HEADING = re.compile(r"^\s*\d{1,2}[.)]\s+(?P<name>\D.{0,80})$")
PAGE_LINE = re.compile(r"^(?:page\s+)?\d+\s*(?:of|/)\s*\d+$", re.I)
NOT_CATEGORY = {"summary", "overview", "contents", "table of contents", "notes", "index"}


def _cell(value) -> str:
    """Table cell as one clean line; wrapped PDF lines are joined ("single-\\nservice" -> "single-service")."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"(?<=\w)-\s*\n\s*(?=\w)", "-", str(value))
    return " ".join(text.split())


def _header_key(cell: str) -> str:
    return " ".join(re.sub(r"[^\w#.]+", " ", cell.lower()).split()).strip(" .")


def _header_roles(row: list[str]) -> dict | None:
    roles: dict[str, int] = {}
    for i, cell in enumerate(row):
        key = _header_key(cell)
        if COUNT_HEADER.match(key):
            roles.setdefault("count", i)
            continue
        for role, names in COLUMN_ROLES.items():
            if key in names:
                roles.setdefault(role, i)
    if "summary" in roles or ("category" in roles and "count" in roles):
        return roles
    return None


def _find_header(rows: list[list[str]]) -> tuple[int, dict | None]:
    for i, row in enumerate(rows[:6]):
        roles = _header_roles(row)
        if roles:
            return i, roles
    return -1, None


def _heading(text: str, meta, body_size: float | None) -> str | None:
    """'title', 'heading' or None for a line outside tables."""
    if not text or len(text) > 90 or PAGE_LINE.match(text) or line_date(text):
        return None
    if meta in ("title", "heading"):
        return meta
    if NUMBERED_HEADING.match(text):
        return "heading"
    if isinstance(meta, (int, float)) and body_size and meta >= body_size * 1.25:
        return "large"
    return None


def parse_blocks(blocks: list[tuple]) -> dict | None:
    """Parse ("line", text, meta) / ("table", rows) blocks. Returns None when no task table exists."""
    body_size = next((b[1] for b in blocks if b[0] == "body_size"), None)
    title = category = None
    seen_table = False
    last = None  # (roles, width) of the previous task table, for tables split across pages
    tasks: list[dict] = []
    expected: dict[str, int] = {}
    expected_total = None
    seen: set = set()
    duplicates = 0

    for block in blocks:
        if block[0] == "body_size":
            continue
        if block[0] == "line":
            level = _heading(block[1], block[2], body_size)
            if not level:
                continue
            numbered = NUMBERED_HEADING.match(block[1])
            name = (numbered.group("name") if numbered else block[1]).strip()
            if not numbered and title is None and not seen_table and level in ("title", "large", "heading"):
                title = name  # the first heading is the document title, not a category
                continue
            category = None if name.lower().rstrip(":") in NOT_CATEGORY else name
            continue

        rows = [[_cell(c) for c in row] for row in block[1]]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        seen_table = True
        start, roles = _find_header(rows)
        if roles is None and last and len(rows[0]) == last[1] and "date" in last[0] \
                and parse_date(rows[0][last[0]["date"]]):
            start, roles = -1, last[0]  # continuation of a table whose header was on the previous page
        if roles is None:
            continue
        get = lambda r, role: r[roles[role]] if role in roles and roles[role] < len(r) else ""

        if "summary" not in roles:  # summary table: "Category | No. of Tasks | ..."
            for r in rows[start + 1:]:
                name, count = get(r, "category"), get(r, "count")
                if not name or not re.fullmatch(r"\d+", count):
                    continue
                if name.lower() in ("total", "grand total"):
                    expected_total = int(count)
                else:
                    expected[name] = int(count)
            continue

        last = (roles, len(rows[start + 1] if start + 1 < len(rows) else rows[0]))
        current = None
        for r in rows[start + 1:]:
            text = get(r, "summary")
            if not text or _header_roles(r):
                continue
            raw_date = get(r, "date")
            when = parse_date(raw_date) or line_date(raw_date) if raw_date else None
            current = when or current  # a blank date cell repeats the row above (merged cells)
            cat = get(r, "category") or category
            key = (current, " ".join(text.lower().split()), (cat or "").lower())
            if key in seen:  # e.g. the same rows on a "by category" and a "by date" sheet
                duplicates += 1
                continue
            seen.add(key)
            tasks.append(make_task(len(tasks) + 1, text, current, category=cat,
                                   status_text=get(r, "status"), kind_text=get(r, "kind")))

    if last is None:
        return None
    if expected and expected_total is None:
        expected_total = sum(expected.values())
    return {
        "tasks": tasks,
        "title": title,
        "layout": "table",
        "expected": {"total": expected_total, "categories": expected} if expected_total is not None else None,
        "duplicates": duplicates,
    }


def _pdf_blocks(data: bytes) -> list[tuple]:
    import pymupdf
    blocks = []
    sizes: Counter = Counter()  # characters per font size, to tell headings from body text
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for page in pdf:
            try:
                tables = page.find_tables().tables
            except Exception:  # table detection is best effort
                tables = []
            items = [(t.bbox[1], ("table", t.extract())) for t in tables]
            boxes = [pymupdf.Rect(t.bbox) for t in tables]
            for b in page.get_text("dict")["blocks"]:
                for line in b.get("lines", []):
                    text = "".join(s["text"] for s in line["spans"]).strip()
                    size = max(s["size"] for s in line["spans"])
                    sizes[round(size, 1)] += len(text)
                    rect = pymupdf.Rect(line["bbox"])
                    if not text or any(box.contains(rect.tl + (rect.br - rect.tl) * 0.5) for box in boxes):
                        continue
                    items.append((rect.y0, ("line", text, size)))
            items.sort(key=lambda item: item[0])
            blocks += [item[1] for item in items]
    if sizes:
        blocks.insert(0, ("body_size", sizes.most_common(1)[0][0]))
    return blocks


def _docx_blocks(data: bytes) -> list[tuple]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = docx.Document(io.BytesIO(data))
    blocks = []
    for el in document.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, document)
            style = (p.style.name or "").lower() if p.style is not None else ""
            level = "title" if style == "title" else "heading" if style.startswith("heading") else None
            blocks.append(("line", p.text.strip(), level))
        elif tag == "tbl":
            blocks.append(("table", [[c.text for c in row.cells] for row in Table(el, document).rows]))
    return blocks


def _xlsx_blocks(data: bytes) -> list[tuple]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [("table", list(ws.iter_rows(values_only=True))) for ws in wb.worksheets]
    finally:
        wb.close()


def _text_blocks(text: str) -> list[tuple]:
    """Markdown-style text: '| a | b |' rows become tables and '#' lines become headings."""
    blocks, table = [], None
    for raw in text.splitlines():
        s = raw.strip()
        if s.startswith("|") and s.count("|") >= 2:
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                table = (table or []) + [cells]
            continue
        if table:
            blocks.append(("table", table))
            table = None
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(("line", m.group(2).strip(), "title" if len(m.group(1)) == 1 else "heading"))
        else:
            blocks.append(("line", s, None))
    if table:
        blocks.append(("table", table))
    return blocks


def _lines_result(text: str) -> dict:
    return {"tasks": parse_tasks(text), "title": None, "layout": "lines", "expected": None, "duplicates": 0}


def parse_document(filename: str, data: bytes) -> dict:
    """Tasks plus document details: {"tasks", "title", "layout", "expected", "duplicates"}."""
    name = filename.lower()
    if name.endswith(".xls"):
        raise ValueError("Legacy .xls files are not supported. Save it as .xlsx and try again.")
    if name.endswith(".pdf"):
        blocks = _pdf_blocks(data)
    elif name.endswith(".docx"):
        blocks = _docx_blocks(data)
    elif name.endswith((".xlsx", ".xlsm")):
        blocks = _xlsx_blocks(data)
        result = parse_blocks(blocks)
        if result:
            return result
        rows = (" | ".join(_cell(c) for c in row if _cell(c)) for b in blocks for row in b[1])
        return _lines_result("\n".join(rows))
    elif name.endswith(".csv"):
        text = extract_text(filename, data)
        return parse_blocks([("table", list(csv.reader(io.StringIO(text))))]) or _lines_result(text)
    else:
        return parse_text(extract_text(filename, data))
    return parse_blocks(blocks) or _lines_result(extract_text(filename, data))


def parse_text(text: str) -> dict:
    """Pasted or plain text: Markdown tables if present, otherwise the line layout."""
    return parse_blocks(_text_blocks(text)) or _lines_result(text)
